import os
import pandas as pd
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                     TableStyle, PageBreak)
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    HAS_REPORTLAB = True
except ImportError:
    HAS_REPORTLAB = False


class ReportGenerator:
    # ================= PDF =================
    @staticmethod
    def export_pdf_report(stats, reg_stats, file_path, period_name="Barchasi",
                           header_text="O'ZBEKISTON RESPUBLIKASI KADASTR AGENTLIGI"):
        """Umumiy PDF hisobot."""
        if not HAS_REPORTLAB:
            return None

        # Shriftni ro'yxatga olish (kirill uchun)
        font_name = "Helvetica"
        try:
            # Windows'da standart shriftlar
            for fp in [
                "C:/Windows/Fonts/arial.ttf",
                "C:/Windows/Fonts/DejaVuSans.ttf",
                "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            ]:
                if os.path.exists(fp):
                    pdfmetrics.registerFont(TTFont('CustomFont', fp))
                    font_name = 'CustomFont'
                    break
        except Exception:
            pass

        doc = SimpleDocTemplate(file_path, pagesize=A4,
                                topMargin=1.5 * cm, bottomMargin=1.5 * cm,
                                leftMargin=2 * cm, rightMargin=1.5 * cm)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle('CustomTitle', parent=styles['Title'],
                                      fontName=font_name, fontSize=14, leading=18,
                                      alignment=1, spaceAfter=10)
        h2_style = ParagraphStyle('CustomH2', parent=styles['Heading2'],
                                   fontName=font_name, fontSize=12, leading=16,
                                   spaceBefore=12, spaceAfter=6)
        body_style = ParagraphStyle('CustomBody', parent=styles['BodyText'],
                                     fontName=font_name, fontSize=10, leading=14)
        small_style = ParagraphStyle('CustomSmall', parent=styles['BodyText'],
                                      fontName=font_name, fontSize=9, leading=12)

        elements = []

        # Header
        elements.append(Paragraph(header_text, title_style))
        davr = f"{datetime.now().year}-yil holatiga ko'ra" if period_name in ["Barchasi", ""] else f"{period_name} holatiga ko'ra"
        elements.append(Paragraph(f"MA'LUMOTNOMA<br/>({davr})", title_style))
        elements.append(Spacer(1, 0.5 * cm))

        # Kirish
        intro = (f"Hisobot davrida jami <b>{stats['total']}</b> ta murojaat kelib tushgan "
                 f"(Telegram bot: <b>{stats['tg_total']}</b> ta, "
                 f"Ishonch telefoni: <b>{stats['phone_total']}</b> ta).")
        elements.append(Paragraph(intro, body_style))
        elements.append(Spacer(1, 0.4 * cm))

        # Jadval
        elements.append(Paragraph("Asosiy ko'rsatkichlar:", h2_style))

        data = [
            ["T/r", "Ko'rsatkich nomi", "Soni (ta)"],
            ["1", "Jami murojaatlar", str(stats['total'])],
            ["2", "Agentlik hududiy xodimlarida o'rganishda", str(stats['agentlik_organish'])],
            ["3", "Palata hududiy xodimlarida o'rganishda", str(stats['palata_organish'])],
            ["4", "O'rganib chiqilgan", str(stats['natija_kiritilgan'])],
            ["5", "Asossiz deb topilgan", str(stats['asossiz'])],
            ["6", "Muddati o'tgan (SLA)", str(stats.get('muddati_otgan', 0))],
            ["7", "Ogohlantirish", str(stats.get('ogohlantirish', 0))],
            ["8", "Takroriy murojaatlar", str(stats.get('takroriy_soni', 0))],
            ["9", "Intizomiy choralar", str(stats.get('chora_krilgan_soni', 0))],
        ]

        tbl = Table(data, colWidths=[1.5 * cm, 10 * cm, 3 * cm])
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F497D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('ALIGN', (1, 1), (1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), font_name),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ]))
        elements.append(tbl)
        elements.append(Spacer(1, 0.5 * cm))

        # Hududlar
        if reg_stats is not None and len(reg_stats) > 0:
            elements.append(Paragraph("Eng ko'p murojaat kelib tushgan hududlar:", h2_style))
            for reg, cnt in reg_stats.head(5).items():
                elements.append(Paragraph(f"• {reg}: <b>{cnt}</b> ta murojaat", body_style))

        elements.append(Spacer(1, 1 * cm))
        elements.append(Paragraph(f"Sana: {datetime.now().strftime('%d.%m.%Y')}", small_style))

        doc.build(elements)
        return file_path

    # ================= WORD (XULOSA) =================
    @staticmethod
    def generate_resolution_report(rec, output_path, header_text):
        if not HAS_DOCX:
            return None
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

        p_h = doc.add_paragraph()
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p_h.add_run(header_text + "\n\n")
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(13)

        r_title = p_h.add_run(f"MA'LUMOTNOMA\n(Murojaat № {rec.get('#')} ijrosi bo'yicha)\n")
        r_title.bold = True
        r_title.font.name = 'Times New Roman'
        r_title.font.size = Pt(14)
        r_title.font.color.rgb = RGBColor(15, 37, 55)

        table = doc.add_table(rows=6, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        rows_data = [
            ("Kelib tushgan sana:", str(rec.get('Yaratilgan sana'))),
            ("Murojaat manbasi:", str(rec.get('Manba', 'Telegram bot'))),
            ("Kategoriya:", str(rec.get('Kategoriya', '—'))),
            ("Fuqaro (F.I.Sh.) va Tel:", f"{rec.get('F.I.Sh.')} ({rec.get('Telefon')})"),
            ("Hudud (Viloyat / Tuman):", f"{rec.get('Viloyat')}, {rec.get('Tuman')}"),
            ("Yakuniy holat / Chora:", f"{rec.get('Ijro_Holati')} / {rec.get('Chora_Turi')}"),
        ]

        for i, (k, v) in enumerate(rows_data):
            c_k = table.rows[i].cells[0]
            c_v = table.rows[i].cells[1]
            c_k.width = Inches(2.2)
            c_v.width = Inches(4.5)
            r_k = c_k.paragraphs[0].add_run(k)
            r_k.bold = True
            r_k.font.name = 'Times New Roman'
            r_k.font.size = Pt(11)
            r_v = c_v.paragraphs[0].add_run(v)
            r_v.font.name = 'Times New Roman'
            r_v.font.size = Pt(11)

        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_before = Pt(12)
        r_m_t = p_m.add_run("Murojaatning qisqacha mazmuni:")
        r_m_t.bold = True
        r_m_t.font.name = 'Times New Roman'
        r_m_t.font.size = Pt(12)

        p_mt = doc.add_paragraph()
        r_mt = p_mt.add_run(str(rec.get('Murojaat matni', '')))
        r_mt.font.name = 'Times New Roman'
        r_mt.font.size = Pt(11)

        p_n = doc.add_paragraph()
        p_n.paragraph_format.space_before = Pt(12)
        r_n_t = p_n.add_run("O'rganish xulosasi va ko'rilgan choralar:")
        r_n_t.bold = True
        r_n_t.font.name = 'Times New Roman'
        r_n_t.font.size = Pt(12)

        p_nt = doc.add_paragraph()
        r_nt = p_nt.add_run(str(rec.get('Organish_Natijasi', '')))
        r_nt.font.name = 'Times New Roman'
        r_nt.font.size = Pt(11)
        r_nt.italic = True

        p_sign = doc.add_paragraph()
        p_sign.paragraph_format.space_before = Pt(30)
        p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_s = p_sign.add_run(f"O'rganish o'tkazgan xodim:\n{str(rec.get('Masul_Komplayens'))} ____________")
        r_s.font.name = 'Times New Roman'
        r_s.font.size = Pt(11)

        doc.save(output_path)
        return output_path

    # ================= EXCEL =================
    @staticmethod
    def export_excel(df, file_path):
        export_df = df.copy()
        if 'DT' in export_df.columns and export_df['DT'].notnull().any():
            export_df = export_df.sort_values(by='DT', ascending=False)
        elif 'Yaratilgan sana' in export_df.columns:
            export_df = export_df.sort_values(by='Yaratilgan sana', ascending=False)

        cols = ['Yaratilgan sana', 'Manba', 'Kategoriya', 'F.I.Sh.', 'Telefon', 'Viloyat',
                'Tuman', 'Yoʻnalish', 'Masul_Komplayens', 'Ijro_Holati',
                'Chora_Turi', 'Organish_Natijasi', 'Murojaat matni']
        available = [c for c in cols if c in export_df.columns]
        sub_df = export_df[available].copy()
        sub_df.insert(0, 'T/r', range(1, len(sub_df) + 1))
        sub_df = sub_df.rename(columns={
            'Yaratilgan sana': 'Kelib tushgan sana',
            'Manba': 'Murojaat manbasi',
            'Kategoriya': 'Kategoriya',
            'Yoʻnalish': 'Yoʻnalish turi',
            'Masul_Komplayens': 'Mas’ul komplayens organi',
            'Ijro_Holati': 'Ijro holati',
            'Chora_Turi': 'Ko‘rilgan chora',
            'Organish_Natijasi': 'O‘rganish natijasi',
            'Murojaat matni': 'Murojaat mazmuni'
        })

        with pd.ExcelWriter(file_path, engine='openpyxl') as writer:
            sub_df.to_excel(writer, sheet_name="Murojaatlar", index=False)
            ws = writer.sheets["Murojaatlar"]
            ws.views.sheetView[0].showGridLines = True

            navy_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
            zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            tb = Border(left=Side(style='thin', color='B0C4DE'),
                        right=Side(style='thin', color='B0C4DE'),
                        top=Side(style='thin', color='B0C4DE'),
                        bottom=Side(style='thin', color='B0C4DE'))

            ws.row_dimensions[1].height = 32
            for col_idx in range(1, len(sub_df.columns) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = navy_fill
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = tb

            for row_idx in range(2, len(sub_df) + 2):
                ws.row_dimensions[row_idx].height = 28
                for col_idx in range(1, len(sub_df.columns) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.font = Font(name="Calibri", size=10)
                    cell.border = tb
                    if row_idx % 2 == 0:
                        cell.fill = zebra_fill
                    cell.alignment = Alignment(vertical="center", wrap_text=True)

            col_widths = {
                'T/r': 8, 'Kelib tushgan sana': 19, 'Murojaat manbasi': 22,
                'Kategoriya': 22, 'F.I.Sh.': 26, 'Telefon': 16, 'Viloyat': 22,
                'Tuman': 20, 'Yoʻnalish turi': 32, 'Mas’ul komplayens organi': 36,
                'Ijro holati': 18, 'Ko‘rilgan chora': 22,
                'O‘rganish natijasi': 45, 'Murojaat mazmuni': 55
            }
            for col_idx, col_name in enumerate(sub_df.columns, 1):
                ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_name, 22)

            ws.auto_filter.ref = f"A1:{get_column_letter(len(sub_df.columns))}{len(sub_df) + 1}"

    # ================= WORD (UMUMIY) =================
    @staticmethod
    def export_word_report(stats, reg_stats, file_path, period_name="Barchasi",
                            header_text="O‘ZBEKISTON RESPUBLIKASI KADASTR AGENTLIGI"):
        if not HAS_DOCX:
            return None
        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(0.8)

        p_head = doc.add_paragraph()
        p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p_head.add_run(header_text + "\n")
        r1.bold = True
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(13)

        davr_matni = (f"{datetime.now().year}-yil holatiga ko‘ra" if period_name in ["Barchasi", ""]
                      else f"{period_name} holatiga ko‘ra")
        r3 = p_head.add_run(f"\nMA’LUMOTNOMA\n({davr_matni})\n")
        r3.bold = True
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(14)
        r3.font.color.rgb = RGBColor(15, 37, 55)

        p_body = doc.add_paragraph()
        p_body.paragraph_format.line_spacing = 1.15
        p_body.paragraph_format.first_line_indent = Inches(0.4)
        r_body = p_body.add_run(
            f"Hisobot davrida jami {stats['total']} ta murojaat kelib tushgan "
            f"(Telegram bot: {stats['tg_total']} ta, Telefon: {stats['phone_total']} ta).")
        r_body.font.name = 'Times New Roman'
        r_body.font.size = Pt(12)

        table = doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        hdr_cells = table.rows[0].cells
        titles = ["T/r", "Ko‘rsatkich nomi", "Soni (ta)"]
        for idx, title in enumerate(titles):
            hdr_cells[idx].text = title
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

        data_rows = [
            ("1", "Jami ko‘rib chiqilayotgan murojaatlar", str(stats['total'])),
            ("2", "Agentlik hududiy komplayens xodimlarida o‘rganishda", str(stats['agentlik_organish'])),
            ("3", "Palata hududiy komplayens xodimlarida o‘rganishda", str(stats['palata_organish'])),
            ("4", "O‘rganib chiqilgan (chora ko‘rilgan / asossiz)", str(stats['natija_kiritilgan'])),
            ("5", "— Shundan Asossiz deb topilganlar", str(stats['asossiz'])),
            ("6", "Muddati o‘tgan (SLA)", str(stats.get('muddati_otgan', 0))),
        ]
        for row in data_rows:
            row_cells = table.add_row().cells
            row_cells[0].text = row[0]
            row_cells[1].text = row[1]
            row_cells[2].text = row[2]
            row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for c in row_cells:
                for r in c.paragraphs[0].runs:
                    r.font.name = 'Times New Roman'
                    r.font.size = Pt(11)

        p_reg = doc.add_paragraph()
        p_reg.paragraph_format.space_before = Pt(12)
        r_reg_h = p_reg.add_run("Eng ko‘p murojaat kelib tushgan hududlar:")
        r_reg_h.bold = True
        r_reg_h.font.name = 'Times New Roman'
        r_reg_h.font.size = Pt(12)

        for reg, cnt in reg_stats.head(5).items():
            p_item = doc.add_paragraph()
            p_item.paragraph_format.left_indent = Inches(0.2)
            r_item = p_item.add_run(f"• {reg}: {cnt} ta murojaat")
            r_item.font.name = 'Times New Roman'
            r_item.font.size = Pt(11)

        doc.save(file_path)
        return file_path
